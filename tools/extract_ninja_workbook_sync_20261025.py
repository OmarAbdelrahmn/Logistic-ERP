"""Read the source workbook without evaluating instructions or exporting passwords."""
import datetime as dt
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r'C:\Users\omarf\Downloads\تسكين نينجا سيستم.xlsx')
OUT = ROOT / 'outputs/ninja-workbook-sync-20261025'
OUT.mkdir(exist_ok=True)
workbook = openpyxl.load_workbook(SOURCE, read_only=False, data_only=False)
sheets = []
for sheet in workbook:
    password_column = 4 if sheet.title == 'البدلاء' else 5
    rows = []
    for cells in sheet:
        values = [cell.value for cell in cells]
        if not any(value is not None for value in values):
            continue
        if cells[0].row > 1:
            password = values[password_column]
            values[password_column] = None if password is None else dict(
                present=True, length=len(str(password)), numeric=isinstance(password, (int, float)),
                numberFormat=cells[password_column].number_format,
            )
        rows.append(dict(row=cells[0].row, values=values))
    sheets.append(dict(sheet=sheet.title, maxRow=sheet.max_row, maxColumn=sheet.max_column,
                       merged=list(map(str, sheet.merged_cells.ranges)), rows=rows))
result = dict(source=str(SOURCE), sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              modifiedUtc=dt.datetime.fromtimestamp(SOURCE.stat().st_mtime, dt.timezone.utc).isoformat(),
              sheets=sheets)
(OUT / 'source-extracted.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps(dict(sha256=result['sha256'], sheets=[dict(name=s['sheet'], nonemptyRows=len(s['rows'])) for s in sheets])))
