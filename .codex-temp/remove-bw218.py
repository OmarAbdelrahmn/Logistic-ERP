import re, json, zipfile, hashlib
from pathlib import Path
from openpyxl import load_workbook

source = Path(r'C:\Users\omarf\Downloads\التسكين.xlsx')
dest = Path('outputs/fuel-card-assignment-20261006/التسكين.xlsx')
dest.parent.mkdir(parents=True, exist_ok=True)
book = load_workbook(source, data_only=False)
assert book.sheetnames == ['ورقة1']
sheet = book.active
assert [(sheet.cell(r,1).value,str(sheet.cell(r,2).value)) for r in [57,58]] == [('BW218','2591951724'),('BW218','2578632750')]
assert not sheet.merged_cells.ranges and not sheet.tables
assert not sheet.data_validations.count and not len(sheet.conditional_formatting)
assert sheet.auto_filter.ref in [None,'A1:C228']
before = [list(row) for row in sheet.iter_rows(values_only=True)]
expected = before[:56] + before[58:]
with zipfile.ZipFile(source) as zin, zipfile.ZipFile(dest,'w') as zout:
    for info in zin.infolist():
        payload = zin.read(info.filename)
        if info.filename == 'xl/worksheets/sheet1.xml':
            xml = payload.decode('utf-8')
            assert '<f' not in xml and '<drawing' not in xml
            def move_row(match):
                row = match.group(0)
                idx = int(re.search(r'<row\b[^>]*\br="(\d+)"',row).group(1))
                if idx in [57,58]: return ''
                if idx > 58:
                    row = re.sub(r'(<row\b[^>]*\br=")\d+(\")',lambda m:m.group(1)+str(idx-2)+m.group(2),row,count=1)
                    row = re.sub(r'(<c\b[^>]*\br=")([A-Z]+)(\d+)(\")',lambda m:m.group(1)+m.group(2)+str(int(m.group(3))-2)+m.group(4),row)
                return row
            xml = re.sub(r'<row\b[^>]*>.*?</row>',move_row,xml,flags=re.S)
            xml = xml.replace('ref="A1:C228"','ref="A1:C226"')
            xml = xml.replace('ref="A1:C684"','ref="A1:C682"')
            payload = xml.encode('utf-8')
        elif info.filename == 'xl/workbook.xml':
            payload = payload.replace(b'$A$1:$C$228',b'$A$1:$C$226')
        zout.writestr(info,payload)
afterbook = load_workbook(dest,data_only=False)
assert [list(row) for row in afterbook.active.iter_rows(values_only=True)] == expected
assert afterbook.active.max_row == sheet.max_row - 2
assert sum(any(v is not None for v in row) for row in expected) == 226
assert all(row[0] != 'BW218' for row in expected)
with zipfile.ZipFile(source) as old, zipfile.ZipFile(dest) as new:
    assert old.namelist()==new.namelist()
    assert all(old.read(n)==new.read(n) for n in old.namelist() if n not in ['xl/worksheets/sheet1.xml','xl/workbook.xml'])
    assert zipfile.ZipFile(dest).testzip() is None
print(json.dumps({'output':str(dest.resolve()),'removedRows':2,'remainingRows':225,'sha256':hashlib.sha256(dest.read_bytes()).hexdigest()},ensure_ascii=False))
