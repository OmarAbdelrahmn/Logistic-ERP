import json
from pathlib import Path
from openpyxl import load_workbook

book = load_workbook(r'C:\Users\omarf\Downloads\التسكين.xlsx', data_only=True, read_only=True)
data = []
for sheet in book:
    rows = [{'row': i, 'values': list(row)} for i, row in enumerate(sheet.iter_rows(values_only=True), 1) if any(v is not None for v in row)]
    data.append({'name': sheet.title, 'rows': rows})
Path(__file__).with_name('fuel-assignment-source.json').write_text(json.dumps(data, ensure_ascii=False, default=str, indent=2), encoding='utf-8')
print(json.dumps(data, ensure_ascii=False, default=str))
