"""Generate the reviewed SQL seed from the supplied Riyadh housing workbook."""

from collections import Counter, defaultdict
from pathlib import Path
import re
import sys

from openpyxl import load_workbook

SOURCE = Path(sys.argv[1])
DESTINATION = Path(__file__).with_name("seed-riyadh-housing-20260929.sql")
TARGET_ID = "01a0d372-4ac7-7e6e-bd47-cbfbe074dafd"
CITY_ID = "019c18d5-62e1-7000-8000-000000000004"
CODES = {
    "الأمير بندر 1": "RYD-PRINCE-BANDAR-1",
    "النسيم 1": "الرياض النسيم",
    "النسيم 2": "RYD-NASEEM-2",
    "النسيم 3": "RYD-NASEEM-3",
}
NAMES = {
    "الأمير بندر 1": "Prince Bandar 1",
    "النسيم 1": "Riyadh Al Naseem 1",
    "النسيم 2": "Riyadh Al Naseem 2",
    "النسيم 3": "Riyadh Al Naseem 3",
}


def literal(value):
    if value is None:
        return "NULL"
    if isinstance(value, int):
        return str(value)
    return "N'" + str(value).replace("'", "''") + "'"


def equipment(value):
    result = []
    for line in str(value or "").splitlines():
        line = line.strip()
        if not line:
            continue
        match = re.match(r"^(\d+)\s*(.+)$", line)
        quantity = int(match.group(1)) if match else 1
        name = match.group(2).strip() if match else line
        if "مرتب" in name or "مراتب" in name:
            name = "مرتبة"
        elif "وساد" in name:
            name = "وسادة"
        elif "تكيف" in name:
            name = "مكيف"
        elif "غسال" in name:
            name = "غسالة"
        elif "تلاج" in name:
            name = "ثلاجة"
        elif "بوتجاز" in name:
            name = "بوتجاز"
        elif "اسطوان" in name:
            name = "اسطوانة غاز"
        result.append((name, quantity))
    return result


sheet = load_workbook(SOURCE, data_only=True).active
rooms = {}
people = []
items = defaultdict(Counter)
totals = {}
for row in sheet.iter_rows(min_row=2):
    place, number, iqama, person, capacity, expected, notes, room_items, floor_items, floor_total = [cell.value for cell in row[:10]]
    if place is None or number is None:
        continue
    place = str(place).strip()
    number = str(number).strip()
    key = (place, number)
    if key not in rooms:
        rooms[key] = [capacity, notes, expected]
    elif capacity is not None:
        rooms[key][0] = capacity
    if notes and not rooms[key][1]:
        rooms[key][1] = notes
    if iqama is not None or person is not None:
        if not person:
            raise ValueError(f"Row {row[0].row}: occupant has no name")
        iqama = str(iqama).strip() if iqama is not None else ""
        if iqama != "خارجى" and not re.fullmatch(r"\d{10}", iqama):
            raise ValueError(f"Row {row[0].row}: invalid iqama {iqama}")
        people.append((place, number, None if iqama == "خارجى" else iqama, str(person).strip(), row[0].row))
    for name, quantity in equipment(room_items):
        items[(place, number)][name] += quantity
    for name, quantity in equipment(floor_items):
        items[(place, None)][name] += quantity
    if floor_total:
        totals[place] = Counter(dict(equipment(floor_total)))

for key, (capacity, _, _) in rooms.items():
    if capacity is None:
        raise ValueError(f"Missing capacity: {key}")
    actual = sum(1 for person in people if person[:2] == key)
    if actual > int(capacity):
        raise ValueError(f"Room over capacity: {key}")

for place in CODES:
    computed = Counter()
    for key, values in items.items():
        if key[0] == place:
            computed.update(values)
    if computed != totals[place]:
        raise ValueError(f"Equipment totals mismatch for {place}: {computed} vs {totals[place]}")

def values(rows):
    return ",\n".join("    (" + ", ".join(literal(value) for value in row) + ")" for row in rows)

housing_values = [(place, code, NAMES[place]) for place, code in CODES.items()]
room_values = [(place, number, int(capacity), notes) for (place, number), (capacity, notes, _) in rooms.items()]
people_values = people
equipment_values = [(place, number, name, quantity) for (place, number), counter in items.items() for name, quantity in counter.items()]

sql = f"""-- Generated from the Riyadh housing workbook. Run after AddHousingPendingOccupants.
-- Replaces the linked Al Naseem 1 room assignments. Existing periods remain as history.
SET XACT_ABORT ON;
SET QUOTED_IDENTIFIER ON;
SET ANSI_NULLS ON;
SET ANSI_PADDING ON;
SET ANSI_WARNINGS ON;
SET ARITHABORT ON;
SET CONCAT_NULL_YIELDS_NULL ON;
SET NUMERIC_ROUNDABORT OFF;
BEGIN TRANSACTION;
IF NOT EXISTS (SELECT 1 FROM [migration].[__ApplicationMigrationsHistory]
               WHERE [MigrationId] LIKE '%AddHousingPendingOccupants')
    THROW 51000, 'Apply the housing migration before seeding.', 1;
IF NOT EXISTS (SELECT 1 FROM [app].[Housing] WHERE [Id] = '{TARGET_ID}'
               AND [CityId] = '{CITY_ID}' AND [IsDeleted] = 0)
    THROW 51001, 'The linked Riyadh Al Naseem housing was not found.', 1;
IF EXISTS (SELECT 1 FROM [app].[HousingRooms] WHERE [HousingId] = '{TARGET_ID}'
           AND [Name] = N'1' AND [IsDeleted] = 0)
    THROW 51002, 'Riyadh housing workbook has already been imported.', 1;

CREATE TABLE #HousingSource ([Place] nvarchar(100) PRIMARY KEY, [Code] nvarchar(32), [NameEn] nvarchar(200));
INSERT INTO #HousingSource VALUES
{values(housing_values)};
CREATE TABLE #RoomSource ([Place] nvarchar(100), [Number] nvarchar(100), [Capacity] int, [Notes] nvarchar(2000));
INSERT INTO #RoomSource VALUES
{values(room_values)};
CREATE TABLE #PersonSource ([Place] nvarchar(100), [Number] nvarchar(100), [Iqama] nvarchar(20), [Name] nvarchar(200), [SourceRow] int);
INSERT INTO #PersonSource VALUES
{values(people_values)};
CREATE TABLE #EquipmentSource ([Place] nvarchar(100), [Number] nvarchar(100), [Name] nvarchar(100), [Quantity] int);
INSERT INTO #EquipmentSource VALUES
{values(equipment_values)};

IF EXISTS (
    SELECT 1 FROM #PersonSource p
    JOIN [app].[Employees] e ON e.[IqamaNo] = p.[Iqama] AND e.[IsDeleted] = 0
    GROUP BY p.[Iqama] HAVING COUNT(e.[Id]) > 1)
    THROW 51003, 'Duplicate employee iqama; no housing was changed.', 1;
IF EXISTS (
    SELECT 1 FROM #PersonSource p
    JOIN [app].[Employees] e ON e.[IqamaNo] = p.[Iqama] AND e.[IsDeleted] = 0
    JOIN [app].[HousingResidencePeriods] x ON x.[EmployeeId] = e.[Id] AND x.[EffectiveTo] IS NULL
    JOIN [app].[HousingRooms] r ON r.[Id] = x.[RoomId]
    WHERE r.[HousingId] <> '{TARGET_ID}')
    THROW 51005, 'A workbook resident is already assigned to another housing.', 1;

-- Close old current residence periods and archive the linked housing's old rooms.
UPDATE p SET [EffectiveTo] = DATEADD(day, -1, CAST(GETDATE() AS date)),
    [MoveOutReason] = N'Replaced by Riyadh housing workbook import'
FROM [app].[HousingResidencePeriods] p
JOIN [app].[HousingRooms] r ON r.[Id] = p.[RoomId]
WHERE r.[HousingId] = '{TARGET_ID}' AND p.[EffectiveTo] IS NULL;
UPDATE r SET [CurrentOccupancy] = 0, [IsDeleted] = 1,
    [DeletionReason] = N'Replaced by Riyadh housing workbook import'
FROM [app].[HousingRooms] r
WHERE r.[HousingId] = '{TARGET_ID}' AND r.[IsDeleted] = 0;

INSERT INTO [app].[Housing] ([Id], [Code], [NameAr], [NameEn], [CityId], [Status], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), s.[Code], s.[Place], s.[NameEn], '{CITY_ID}', 2, SYSUTCDATETIME(), 0
FROM #HousingSource s
WHERE s.[Place] <> N'النسيم 1'
  AND NOT EXISTS (SELECT 1 FROM [app].[Housing] h WHERE h.[Code] = s.[Code]);

INSERT INTO [app].[HousingWarehouses] ([Id], [HousingId], [NameAr], [NameEn], [IsDefault], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), h.[Id], N'المستودع الافتراضي', N'Default warehouse', 1, SYSUTCDATETIME(), 0
FROM #HousingSource s JOIN [app].[Housing] h ON h.[Code] = s.[Code]
WHERE NOT EXISTS (SELECT 1 FROM [app].[HousingWarehouses] w WHERE w.[HousingId] = h.[Id] AND w.[IsDeleted] = 0);

INSERT INTO [app].[HousingFloors] ([Id], [HousingId], [Name], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), h.[Id], N'1', SYSUTCDATETIME(), 0
FROM #HousingSource s JOIN [app].[Housing] h ON h.[Code] = s.[Code]
WHERE NOT EXISTS (SELECT 1 FROM [app].[HousingFloors] f WHERE f.[HousingId] = h.[Id] AND f.[Name] = N'1' AND f.[IsDeleted] = 0);

INSERT INTO [app].[HousingRooms] ([Id], [HousingId], [FloorId], [Name], [Capacity], [CurrentOccupancy], [Notes], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), h.[Id], f.[Id], s.[Number], s.[Capacity], 0, s.[Notes], SYSUTCDATETIME(), 0
FROM #RoomSource s
JOIN #HousingSource hs ON hs.[Place] = s.[Place]
JOIN [app].[Housing] h ON h.[Code] = hs.[Code]
JOIN [app].[HousingFloors] f ON f.[HousingId] = h.[Id] AND f.[Name] = N'1' AND f.[IsDeleted] = 0
WHERE NOT EXISTS (SELECT 1 FROM [app].[HousingRooms] r WHERE r.[HousingId] = h.[Id] AND r.[Name] = s.[Number] AND r.[IsDeleted] = 0);

-- Existing people are assigned by iqama. Name-only people are kept separate from Employees.
INSERT INTO [app].[HousingResidencePeriods]
    ([Id], [EmployeeId], [RoomId], [EffectiveFrom], [AssignedByUserId], [CreatedAtUtc], [SourceReference], [CapacityOverrideUsed])
SELECT NEWID(), e.[Id], r.[Id], CAST(GETDATE() AS date), '00000000-0000-0000-0000-000000000000', SYSUTCDATETIME(), N'Riyadh housing workbook', 0
FROM #PersonSource p
JOIN [app].[Employees] e ON e.[IqamaNo] = p.[Iqama] AND e.[IsDeleted] = 0
JOIN #HousingSource hs ON hs.[Place] = p.[Place]
JOIN [app].[Housing] h ON h.[Code] = hs.[Code]
JOIN [app].[HousingRooms] r ON r.[HousingId] = h.[Id] AND r.[Name] = p.[Number] AND r.[IsDeleted] = 0
WHERE p.[Iqama] IS NOT NULL
  AND NOT EXISTS (SELECT 1 FROM [app].[HousingResidencePeriods] x WHERE x.[EmployeeId] = e.[Id] AND x.[EffectiveTo] IS NULL);

INSERT INTO [app].[HousingExternalOccupants] ([Id], [RoomId], [Name], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), r.[Id], p.[Name], SYSUTCDATETIME(), 0
FROM #PersonSource p
JOIN #HousingSource hs ON hs.[Place] = p.[Place]
JOIN [app].[Housing] h ON h.[Code] = hs.[Code]
JOIN [app].[HousingRooms] r ON r.[HousingId] = h.[Id] AND r.[Name] = p.[Number] AND r.[IsDeleted] = 0
WHERE p.[Iqama] IS NULL;

INSERT INTO [app].[HousingPendingOccupants] ([Id], [RoomId], [IqamaNo], [Name], [SourceRow], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), r.[Id], p.[Iqama], p.[Name], p.[SourceRow], SYSUTCDATETIME(), 0
FROM #PersonSource p
JOIN #HousingSource hs ON hs.[Place] = p.[Place]
JOIN [app].[Housing] h ON h.[Code] = hs.[Code]
JOIN [app].[HousingRooms] r ON r.[HousingId] = h.[Id] AND r.[Name] = p.[Number] AND r.[IsDeleted] = 0
WHERE p.[Iqama] IS NOT NULL AND NOT EXISTS
    (SELECT 1 FROM [app].[Employees] e WHERE e.[IqamaNo] = p.[Iqama] AND e.[IsDeleted] = 0);

INSERT INTO [app].[HousingEquipment] ([Id], [FloorId], [RoomId], [Name], [Quantity], [CreatedAtUtc], [IsDeleted])
SELECT NEWID(), CASE WHEN s.[Number] IS NULL THEN f.[Id] END,
       CASE WHEN s.[Number] IS NOT NULL THEN r.[Id] END,
       s.[Name], s.[Quantity], SYSUTCDATETIME(), 0
FROM #EquipmentSource s
JOIN #HousingSource hs ON hs.[Place] = s.[Place]
JOIN [app].[Housing] h ON h.[Code] = hs.[Code]
JOIN [app].[HousingFloors] f ON f.[HousingId] = h.[Id] AND f.[Name] = N'1' AND f.[IsDeleted] = 0
LEFT JOIN [app].[HousingRooms] r ON r.[HousingId] = h.[Id] AND r.[Name] = s.[Number] AND r.[IsDeleted] = 0;

UPDATE r SET [CurrentOccupancy] = x.[Occupants]
FROM [app].[HousingRooms] r
CROSS APPLY (SELECT
    (SELECT COUNT(*) FROM [app].[HousingResidencePeriods] p WHERE p.[RoomId] = r.[Id] AND p.[EffectiveTo] IS NULL) +
    (SELECT COUNT(*) FROM [app].[HousingExternalOccupants] e WHERE e.[RoomId] = r.[Id] AND e.[IsDeleted] = 0) +
    (SELECT COUNT(*) FROM [app].[HousingPendingOccupants] u WHERE u.[RoomId] = r.[Id] AND u.[IsDeleted] = 0) AS [Occupants]) x
WHERE r.[IsDeleted] = 0 AND r.[HousingId] IN (SELECT h.[Id] FROM [app].[Housing] h JOIN #HousingSource s ON s.[Code] = h.[Code]);

IF EXISTS (SELECT 1 FROM [app].[HousingRooms] r WHERE r.[CurrentOccupancy] > r.[Capacity])
    THROW 51004, 'Imported room exceeds capacity.', 1;
COMMIT TRANSACTION;
SELECT h.[NameAr], COUNT(r.[Id]) AS [Rooms], SUM(r.[CurrentOccupancy]) AS [Occupants]
FROM #HousingSource s JOIN [app].[Housing] h ON h.[Code] = s.[Code]
JOIN [app].[HousingRooms] r ON r.[HousingId] = h.[Id] AND r.[IsDeleted] = 0
GROUP BY h.[NameAr];
SELECT p.[SourceRow], p.[Iqama], p.[Name] AS [UnmatchedName], p.[Place], p.[Number]
FROM #PersonSource p
WHERE p.[Iqama] IS NOT NULL AND NOT EXISTS
    (SELECT 1 FROM [app].[Employees] e WHERE e.[IqamaNo] = p.[Iqama] AND e.[IsDeleted] = 0);
"""
DESTINATION.write_text(sql, encoding="utf-8-sig")
print(f"{len(CODES)} housings, {len(rooms)} rooms, {len(people)} people, {len(equipment_values)} equipment lines")
print("Workbook expected occupants vs named rows:")
for key, (_, _, expected) in rooms.items():
    actual = sum(1 for person in people if person[:2] == key)
    if expected is not None and int(expected) != actual:
        print(f"  {key}: expected {expected}, named {actual}")
