import json, hashlib
from pathlib import Path
from datetime import datetime, date
import openpyxl

p=Path(r'C:\Users\omarf\Downloads\تقرير سيستم  داتا بدلاء .xlsx')
w=openpyxl.load_workbook(p, data_only=False)
out=[]
for s in w:
    rows=[]
    for row in s:
        cells={c.coordinate: (c.value.isoformat() if isinstance(c.value,(datetime,date)) else c.value) for c in row if c.value is not None}
        if cells: rows.append(cells)
    out.append({'sheet':s.title,'rows':s.max_row,'cols':s.max_column,'cells':rows})
Path('tmp/chefz-assignment-workbook-20261013.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print('SHA256',hashlib.sha256(p.read_bytes()).hexdigest())
for s in out:
    print(json.dumps({**s,'cells':s['cells'][:10]},ensure_ascii=False))
