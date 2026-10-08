import csv
from collections import Counter
from pathlib import Path
import openpyxl

source = Path('outputs/iqama-platform-review-20261007-v2/platform_accounts_deduplicated.xlsx')
target = Path('tmp/platform_upload/accounts.csv')
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
platforms = {'جاهز - Jahez': 'JAHEZ', 'هنقرستيشن - HungerStation': 'HUNGER', 'ذا شيفز - The Chefsz': 'CHEFSZ'}
rows = []
names = {}
for sheet_name, platform in platforms.items():
    for row_no, values in enumerate(workbook[sheet_name].iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None for v in values):
            continue
        external, owner_or_name, vehicle, city, sponsor, status = values[:6]
        owner = external if platform == 'CHEFSZ' else owner_or_name
        if platform == 'CHEFSZ' and owner_or_name:
            names[str(owner).strip()] = str(owner_or_name).strip()
        if external is None or owner is None or city is None or sponsor is None or status is None:
            raise ValueError(f'Missing required value: {sheet_name}:{row_no}')
        rows.append({'platform': platform, 'external_id': str(external).strip(), 'owner_iqama': str(owner).strip(),
                     'vehicle': str(vehicle or '').strip(), 'city': str(city).strip(), 'sponsor': str(sponsor).strip(),
                     'status': str(status).strip(), 'sheet_row': row_no})
with target.open('w', encoding='utf-8-sig', newline='') as stream:
    writer = csv.DictWriter(stream, fieldnames=rows[0].keys())
    writer.writeheader()
    writer.writerows(rows)
print('ROWS', len(rows), 'OWNERS', len(set(r['owner_iqama'] for r in rows)))
hunger_book = openpyxl.load_workbook(r'C:\Users\omarf\Downloads\تمام تطبيق هنقرستيشن.xlsx', read_only=True, data_only=True)
for values in hunger_book.active.iter_rows(min_row=2, values_only=True):
    if values[0] is not None and values[1] is not None:
        names.setdefault(str(values[0]).strip(), str(values[1]).strip())
with Path('tmp/platform_upload/owner_names.csv').open('w', encoding='utf-8-sig', newline='') as stream:
    writer = csv.writer(stream)
    writer.writerow(['owner_iqama', 'owner_name'])
    writer.writerows(sorted(names.items()))
for col in ('platform', 'city', 'sponsor', 'status'):
    print(col, dict(Counter(row[col] for row in rows)))
