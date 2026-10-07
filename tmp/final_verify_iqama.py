from openpyxl import load_workbook
from collections import Counter

src=load_workbook(r'C:\Users\omarf\Downloads\دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx',read_only=True,data_only=True)
out=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx',read_only=True,data_only=True)
review=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_upload_review.xlsx',read_only=True,data_only=True)
source_keeta=list(src['كيتا - Keeta'].values)
output_keeta=list(out['كيتا - Keeta'].values)
assert len(source_keeta)==len(output_keeta)
assert all(tuple(a)==tuple(b)+(None,)*(len(a)-len(b)) for a,b in zip(source_keeta,output_keeta))
sets=[]
for p in list(out)[1:]:
    column=4 if p.title.startswith('كيتا') else 0 if p.title.startswith('ذا شيفز') else 1
    start=4 if p.title.startswith('كيتا') else 1
    sets.append({str(row[column]) for row in list(p.values)[start:] if row[0] is not None})
all_ids=set().union(*sets)
categories=Counter(sum(i in x for x in sets) for i in all_ids)
review_rows=list(review.active.values)[5:]
issues=Counter(row[5] for row in review_rows if row[5])
print('unique',len(all_ids),'categories',dict(categories),'issues',dict(issues),'review_rows',len(review_rows))
assert len(all_ids)==524
assert categories==Counter({4:200,3:146,2:125,1:53})
assert issues['تصحيح إقامة']==1
assert issues['حساب مكرر بالكامل']==10
assert issues['حساب مكرر ببيانات مختلفة']==5
assert issues['غير موجودة بسجل المقيمين 2026-09-12']==33
assert len(review_rows)==49
