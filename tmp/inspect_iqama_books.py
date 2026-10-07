from pathlib import Path
from openpyxl import load_workbook
import sys, re

def show(v):
    if v is None: return ''
    s = str(v).strip().replace('\n',' ')
    if re.search(r'\d{7,}', s):
        s = re.sub(r'\d{7,}', lambda m: f'<digits:{len(m.group())}>', s)
    return s[:75]

for name in sys.argv[1:]:
    p=Path(name)
    print('\nFILE',p.name)
    wb=load_workbook(p,read_only=True,data_only=False)
    for ws in wb:
        print('SHEET',ws.title,'rows',ws.max_row,'cols',ws.max_column)
        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,6),max_col=min(ws.max_column,40),values_only=True):
            print(' | '.join(f'{i+1}:{show(v)}' for i,v in enumerate(row) if v is not None))
