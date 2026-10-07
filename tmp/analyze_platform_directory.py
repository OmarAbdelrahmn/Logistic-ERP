from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook
import re

p=Path(r'C:\Users\omarf\Downloads\دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx')
w=load_workbook(p,read_only=True,data_only=False)
for s in w:
    print('\nSHEET',s.title)
    for n,row in enumerate(s.iter_rows(min_row=1,max_row=min(10,s.max_row),max_col=min(18,s.max_column)),1):
        print(n,' '.join(f'{c.column_letter}={str(c.value)[:50] if not re.search(r"\d{7,}",str(c.value)) else "<id>"}' for c in row if c.value is not None))
    if s.title.startswith('لوحة'): continue
    hdr=None
    for row in s.iter_rows(min_row=1,max_row=10,values_only=True):
        if any(isinstance(v,str) and ('اقامة صاحب اليوزر' in v or 'رقم بطاقة الهوية' in v) for v in row):
            hdr=row; break
    if hdr is None: continue
    col=[i for i,v in enumerate(hdr) if isinstance(v,str) and ('اقامة صاحب اليوزر' in v or 'رقم بطاقة الهوية' in v)][0]
    start=4 if s.title.startswith('كيتا') else 2
    ids=[]; blanks=[]; invalid=[]; types=Counter(); values=Counter()
    for row in s.iter_rows(min_row=start,max_row=s.max_row):
        c=row[col];v=c.value
        if not any(x.value is not None for x in row): continue
        if v is None or str(v).strip()=='':blanks.append(c.row);continue
        st=str(v).strip(); ids.append((c.row,st));types[c.data_type]+=1
        if not re.fullmatch(r'[12]\d{9}',st):invalid.append((c.row,st[:3]+'…'+st[-2:] if len(st)>5 else st))
        values[st]+=1
    print('IQAMA', 'rows',len(ids),'blank',len(blanks),blanks[:30],'invalid',len(invalid),invalid[:30],'types',types,'duplicate identities',sum(v-1 for v in values.values() if v>1))
    ac_col=0
    accounts=defaultdict(list)
    for row in s.iter_rows(min_row=start,max_row=s.max_row):
        ac=row[ac_col].value
        if ac is not None: accounts[str(ac).strip()].append(row[0].row)
    print('ACCOUNT',len(accounts),'dups',[(k if len(k)<8 else k[:3]+'…'+k[-2:],v) for k,v in accounts.items() if len(v)>1][:20])
