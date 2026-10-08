from collections import defaultdict,Counter
from openpyxl import load_workbook

master=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx',read_only=True,data_only=True)
new=load_workbook(r'C:\Users\omarf\Downloads\تمام تطبيق هنقرستيشن.xlsx',read_only=True,data_only=True).active
old=master['هنقرستيشن - HungerStation']
status_map={'يعمل':'يعمل (نشط)','ارشيف':'غير نشط (أرشيف)','مشاكل':'موقوف (مشاكل نظامية)'}
sponsor_map={'اكسبرس':'شركة اكسبرس جايت','البوابا':'شركة البوابا الموكبلا','التجارية':'مؤسسة البوابة التجارية','خارج الكفالة':'خارج الكفالة'}
byid=defaultdict(list)
for row in old.iter_rows(min_row=2,values_only=False):
 v=[x.value for x in row];byid[str(v[0])].append((row[0].row,v))
chosen=[];problems=[]
for row in new.iter_rows(min_row=2,values_only=False):
 v=[x.value for x in row];account=str(v[2]);candidates=byid[account]
 matching=[(n,o) for n,o in candidates if str(o[1])==str(v[0]) and o[2]==v[4] and o[3]==v[3] and o[4]==sponsor_map.get(v[5]) and o[5]==status_map.get(v[6])]
 if len(matching)==0:problems.append((account,row[0].row,[(n,o[2:]) for n,o in candidates]))
 else:chosen.append((account,row[0].row,matching[0][0]))
print('selected',len(chosen),'problems',len(problems),problems[:10])
print('selected duplicate accounts')
for a,n,oldrow in chosen:
 if len(byid[a])>1:print(a,'new row',n,'old rows',[r for r,_ in byid[a]],'selected old',oldrow,'statuses',[(r,v[5]) for r,v in byid[a]])
assert len(chosen)==438 and not problems
