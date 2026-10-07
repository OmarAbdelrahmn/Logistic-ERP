from collections import Counter,defaultdict
from openpyxl import load_workbook
from pathlib import Path
import re

base=Path(r'C:\Users\omarf\Downloads')
main=load_workbook(base/'دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx',read_only=True,data_only=True)
ref=load_workbook(base/'المقيمين_ERP_مكتمل_السجلات_الثلاثة_12-09-2026.xlsx',read_only=True,data_only=True)
roster=load_workbook(base/'Riders_List_2026-08-27 (3).xlsx',read_only=True,data_only=True)
housing=load_workbook(base/'بيانات_السكن_الموحدة_ERP_جدة.xlsx',read_only=True,data_only=True)

def norm(v):
    if v is None:return ''
    s=str(v).strip()
    if s.endswith('.0'):s=s[:-2]
    return re.sub(r'[^0-9]','',s)
def mask(v):
    v=norm(v)
    return v[:3]+'…'+v[-2:] if len(v)>=6 else v
def vals(w,sheet,start,col):
    s=w[sheet];d=defaultdict(list)
    for row in s.iter_rows(min_row=start,values_only=False):
        if len(row)>col:
            v=norm(row[col].value)
            if v:d[v].append(row[0].row if hasattr(row[0],'row') else None)
    return d
refs={
 'residents':vals(ref,'المقيمين النشطين',2,0),
 'roster':vals(roster,'Riders',2,0),
 'housing':vals(housing,'بيانات الرفع لنظام ERP',2,0),
}
allrows=[]
for s in list(main)[1:]:
    start=5 if s.title.startswith('كيتا') else 2
    col=4 if s.title.startswith('كيتا') else 1 if not s.title.startswith('ذا شيفز') else 0
    d=defaultdict(list); accounts=defaultdict(list); allids=[]
    for row in s.iter_rows(min_row=start,values_only=False):
        iq=norm(row[col].value)
        if not iq:continue
        n=row[0].row
        acc=str(row[0].value).strip() if row[0].value is not None else ''
        d[iq].append(n);accounts[acc].append((n,iq))
        allrows.append((s.title,n,iq,acc))
        allids.append(iq)
    print('\n',s.title,'rows',len(allids),'unique iqama',len(d))
    for k,refd in refs.items():
        hit=[v for v in allids if v in refd]
        print(k,'row matches',len(hit),'unique',len(set(hit)))
    print('same account multi iqama',[(a,[(n,mask(i)) for n,i in r]) for a,r in accounts.items() if len(set(i for n,i in r))>1])
    print('same iqama multi account',[(mask(i),r) for i,r in d.items() if len(r)>1][:30])
    print('leading digits',Counter(i[0] for i in allids))
ids=defaultdict(list)
for s,n,i,a in allrows:ids[i].append((s,n))
print('\nTOTAL rows',len(allrows),'unique',len(ids),'reference combined',len(set().union(*[set(x) for x in refs.values()])))
print('unmatched all refs',sum(1 for i in ids if not any(i in d for d in refs.values())))
print('matched residents',sum(1 for i in ids if i in refs['residents']))
print('matched roster',sum(1 for i in ids if i in refs['roster']))
print('matched housing',sum(1 for i in ids if i in refs['housing']))
for s,n,i,a in allrows:
    if not re.fullmatch(r'[12]\d{9}',i): print('INVALID',s,n,mask(i))
