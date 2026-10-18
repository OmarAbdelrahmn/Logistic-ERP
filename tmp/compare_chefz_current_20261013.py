import json, re, collections, hashlib, csv
from pathlib import Path
from datetime import datetime, date, timedelta
import openpyxl

ROOT=Path.cwd()
SOURCE=Path('C:/Users/omarf/Downloads/تقرير سيستم  داتا بدلاء .xlsx')
db=json.loads((ROOT/'tmp/chefz-comparison-current-20261013.json').read_text(encoding='utf-8-sig'))
old=json.loads((ROOT/'outputs/chefz-rider-sheet-20261013T112140Z-applied/plan.json').read_text(encoding='utf-8-sig'))
wb=openpyxl.load_workbook(SOURCE,data_only=True,read_only=True)
print('Workbook',[(s.title,s.max_row,s.max_column) for s in wb])
print('Headers',list(next(wb.worksheets[0].iter_rows(values_only=True))))
def norm(v):
    if isinstance(v,float) and v.is_integer(): v=int(v)
    return re.sub(r'\s+','',str(v if v is not None else ''))
def dt(v):
    if isinstance(v,(date,datetime)):return v.strftime('%Y-%m-%d')
    if not v:return ''
    for fmt in ('%d/%m/%Y','%Y-%m-%dT%H:%M:%S','%Y-%m-%d'):
        try:return datetime.strptime(str(v),fmt).date().isoformat()
        except ValueError:pass
    return 'Invalid: '+str(v)
accounts=collections.defaultdict(list)
for a in db['accounts']:accounts[norm(a['ExternalAccountId'])].append(a)
current=collections.defaultdict(list)
for a in db['assignments']:
    if a['Code'] in ('SHIFTZ','CHEFZ') and not a['IsDeleted'] and a['EffectiveTo'] is None:
        current[a['PlatformRiderAccountId']].append(a)
employees=collections.defaultdict(list)
for e in db['employees']:
    if not e['IsDeleted']:employees[norm(e['IqamaNo'])].append(e)
rows=[]
for n,values in enumerate(wb.worksheets[0].iter_rows(min_row=2,values_only=True),2):
    if not any(v is not None for v in values):continue
    rows.append(dict(sheetRow=n,ownerIqama=norm(values[0]),sheetRider=norm(values[1]),sheetName=str(values[2] or ''),sheetStart=dt(values[3])))
duplicates=collections.Counter(r['sheetRider'] for r in rows if r['sheetRider'])
oldrows={r['row']:r for r in old['rows']}
reasons=collections.Counter()
for r in rows:
    aa=accounts[r['ownerIqama']]
    active=[x for a in aa for x in current[a['Id']]]
    r.update(systemRider='; '.join(str(a['RiderIqama'] or '') for a in active),systemStart='; '.join(dt(a['EffectiveFrom']) for a in active),systemAccountStatus='; '.join(str(a['Status']) for a in aa),systemAssignmentStatus='; '.join(str(a['Status']) for a in active))
    if not aa:r['result']='Account missing in system'
    elif len(aa)>1 or len(active)>1:r['result']='Ambiguous system records'
    elif not active:r['result']='Not assigned in system'
    elif norm(active[0]['RiderIqama'])!=r['sheetRider']:r['result']='Different rider'
    elif dt(active[0]['EffectiveFrom'])!=r['sheetStart']:r['result']='Different start date'
    elif active[0]['Status']!=2:r['result']='Assignment status differs'
    else:r['result']='Match'
    why=[]
    if r['result']!='Match':
        if len(active)==1 and r['systemRider']!=r['sheetRider']:why.append('Different assigned rider')
        if len(active)==1 and r['systemStart']!=r['sheetStart']:why.append('Different handover date')
        if not aa:why.append('Chefz account missing')
        elif any(a['Status'] in (3,4,5) for a in aa):why.append('Chefz account unavailable')
        if not r['sheetRider']:why.append('Rider iqama missing in sheet')
        elif duplicates[r['sheetRider']]>1:why.append('Duplicate rider iqama in sheet')
        ee=employees[r['sheetRider']]
        if r['sheetRider'] and len(ee)!=1:why.append('Rider employee missing or ambiguous')
        if len(ee)==1:
            e=ee[0]
            if not e['RiderProfileId'] or e['RiderDeleted']:why.append('Rider profile missing')
            if e['Status']!=3 or e['IsEmployee']:why.append('Rider unavailable')
            other=[a for a in db['assignments'] if a['RiderProfileId']==e['RiderProfileId'] and a['EffectiveTo'] is None and a['Code'] not in ('SHIFTZ','CHEFZ')]
            if other:why.append('Rider assigned to another platform: '+', '.join(sorted(set(a['Code'] for a in other))))
        if not r['sheetStart'] or r['sheetStart'].startswith('Invalid'):why.append('Missing or invalid handover date')
        elif r['sheetStart']>'2026-10-13':why.append('Future handover date')
    r['reasons']='; '.join(why)
    reasons.update(why)
    previous=oldrows.get(r['sheetRow'])
    r['priorImportReasons']='; '.join(previous.get('finalIssues',[])) if previous and previous['ownerIqama']==r['ownerIqama'] and previous['userIqama']==r['sheetRider'] else ''
totals=collections.Counter(r['result'] for r in rows)
sheetowners={r['ownerIqama'] for r in rows}
extra=[a for a in current.values() for a in a if norm(a['ExternalAccountId']) not in sheetowners]
summary={'source':str(SOURCE),'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'readAtUtc':db['capture'][0]['ReadAtUtc'],'sourceRows':len(rows),'sourceUniqueOwners':len(sheetowners),'chefzAccounts':len(db['accounts']),'activeChefzAssignments':sum(map(len,current.values())),'results':dict(totals),'reasons':dict(reasons),'extraAssignedAccounts':extra,'duplicateRiders':{k:v for k,v in duplicates.items() if v>1},'rows':rows}
(ROOT/'tmp/chefz-comparison-analysis-20261013.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in summary.items() if k not in ('rows','extraAssignedAccounts','duplicateRiders')},ensure_ascii=False,indent=2))
print('Extra assigned',len(extra),'Duplicate riders',len(summary['duplicateRiders']),'Duplicate affected rows',sum(summary['duplicateRiders'].values()))
print('Missing account rows',[(r['sheetRow'],r['ownerIqama'],r['sheetRider']) for r in rows if r['result']=='Account missing in system'])
print('Actual rider/date differences',[(r['sheetRow'],r['ownerIqama'],r['result']) for r in rows if r['result'] not in ('Match','Not assigned in system','Account missing in system')])
assert sum(totals.values())==len(rows)
assert len({r['sheetRow'] for r in rows})==len(rows)
out=ROOT/'outputs/chefz-assignment-comparison-20261013'
out.mkdir(parents=True,exist_ok=True)
diff=[r for r in rows if r['result']!='Match']
headers=['sheetRow','ownerIqama','sheetRider','sheetName','sheetStart','systemRider','systemStart','result','reasons','priorImportReasons']
with (out/'differences.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=headers,extrasaction='ignore');writer.writeheader();writer.writerows(diff)
localRead=datetime.fromisoformat(db['capture'][0]['ReadAtUtc'])+timedelta(hours=3)
lines=['# Chefz account assignment comparison','',f'Compared the current ERP database on {localRead:%d October %Y, %H:%M} Riyadh time with `{SOURCE.name}`, worksheet `ورقة1`, rows 2–306.','',
'Match key: column A (owner iqama / Chefz external account ID). Assigned rider: column B (rider iqama). Handover date: column D, interpreted as day/month/year. Current assignments have no end date and are not deleted. Names and IBANs are outside this assignment comparison.','',
'| Result | Rows |','| --- | ---: |']
lines += [f'| {k} | {v} |' for k,v in totals.items()]
lines += [f'| Active Chefz assignments outside the sheet | {len(extra)} |','',f'The system contains {len(db["accounts"])} Chefz accounts. {sum(bool(accounts[r["ownerIqama"]]) for r in rows)} of the {len(rows)} sheet account IDs exist in the system. {len(db["accounts"])-len({a["Id"] for r in rows for a in accounts[r["ownerIqama"]]})} system accounts are outside the supplied sheet; this alone does not indicate a missing assignment.','',
'## Reasons for differences','', 'Reasons can overlap on the same row. These are checked against the current data. Previous import reasons are included in the CSV separately for reference.','', '| Reason | Affected rows |','| --- | ---: |']
lines += [f'| {k} | {v} |' for k,v in reasons.most_common()]
lines += ['', '## All assignment differences','', '| Sheet row | Account owner iqama | Sheet rider iqama | Sheet rider name | Sheet start date | Current rider iqama | Current start date | Difference | Reason |','| ---: | --- | --- | --- | --- | --- | --- | --- | --- |']
for r in diff:
    vals=[r['sheetRow'],r['ownerIqama'],r['sheetRider'] or 'Missing',r['sheetName'],r['sheetStart'] or 'Missing',r['systemRider'] or 'None',r['systemStart'] or 'None',r['result'],r['reasons']]
    lines.append('| '+' | '.join(str(x).replace('|','/').replace('\n',' ') for x in vals)+' |')
lines += ['','## Riders repeated in the sheet','','| Rider iqama | Source rows | Account owner iqamas |','| --- | --- | --- |']
for rider,count in summary['duplicateRiders'].items():
    rr=[r for r in rows if r['sheetRider']==rider]
    lines.append(f'| {rider} | '+', '.join(str(r['sheetRow']) for r in rr)+' | '+', '.join(r['ownerIqama'] for r in rr)+' |')
(out/'comparison-report.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
print('Report directory',out)
