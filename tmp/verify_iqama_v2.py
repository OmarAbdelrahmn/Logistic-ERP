from collections import Counter
from openpyxl import load_workbook
from pathlib import Path
import re

root=Path(r'outputs\iqama-platform-review-20261007-v2')
before=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx',data_only=True)
after=load_workbook(root/'platform_accounts_deduplicated.xlsx',data_only=True)
report=load_workbook(root/'platform_accounts_upload_review.xlsx',data_only=True)
hsource=load_workbook(r'C:\Users\omarf\Downloads\تمام تطبيق هنقرستيشن.xlsx',data_only=True).active
assert before.sheetnames==after.sheetnames
for sheet in ['كيتا - Keeta','ذا شيفز - The Chefsz']:
    assert list(before[sheet].values)==list(after[sheet].values),sheet

j=[r for r in list(after['جاهز - Jahez'].values)[1:] if r[0] is not None]
h=[r for r in list(after['هنقرستيشن - HungerStation'].values)[1:] if r[0] is not None]
c=[r for r in list(after['ذا شيفز - The Chefsz'].values)[1:] if r[0] is not None]
assert (len(j),len(h),len(c))==(380,438,464)
assert len({str(r[0]) for r in j})==380
assert len({str(r[0]) for r in h})==438
assert len({str(r[0]) for r in c})==464
for rows,col in [(j,1),(h,1),(c,0)]:
    assert all(len(r)>=6 and all(v is not None and str(v).strip() for v in r[:6]) for r in rows)
    assert all(re.fullmatch(r'[12]\d{9}',str(r[col])) for r in rows)

seen=set();expected_j=[]
for r in list(before['جاهز - Jahez'].values)[1:]:
    if r[0] is None or str(r[0]) in seen:continue
    seen.add(str(r[0]));expected_j.append(r)
assert j==expected_j
status_map={'يعمل':'يعمل (نشط)','مشاكل':'موقوف (مشاكل نظامية)','ارشيف':'غير نشط (أرشيف)'}
sponsor_map={'اكسبرس':'شركة اكسبرس جايت','البوابا':'شركة البوابا الموكبلا','التجارية':'مؤسسة البوابة التجارية','خارج الكفالة':'خارج الكفالة'}
expected_h=[]
for r in list(hsource.values)[1:]:
    if r[2] is not None: expected_h.append((str(r[2]),str(r[0]),r[4],r[3],sponsor_map[r[5]],status_map[r[6]]))
assert h==expected_h
assert next(r for r in h if str(r[0])=='1305009')[5]=='غير نشط (أرشيف)'
summary=after.worksheets[0]
assert [summary[x].value for x in ['C10','C11','C12','C13','C21']]==[380,438,464,1541,524]
assert [summary[x].value for x in ['D10','E10','D11','E11','D13','E13']]==[309,71,319,119,1270,271]
assert [summary[x].value for x in ['H10','I10','H11','I11','H13']]==[311,69,358,80,1329]
issues=Counter(row[5] for row in list(report.active.values)[5:] if row[5])
assert issues['تم حل تكرار الحساب']==15
assert issues['رقم الإقامة غير موجود في السجل المرجعي']==33
assert issues['فرق في ملخص المدينة']==1
assert issues['فرق في ملخص السجل']==1
assert sum(issues.values())==51
errors=[(s.title,c.coordinate) for s in after for row in s for c in row if c.data_type=='e']
assert not errors,errors[:5]
print('verified: Jahez 380, HungerStation 438, Chefsz 464; 0 duplicate account IDs; 15 resolved review items; 33 historical iqamas still need source verification')
