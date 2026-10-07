from pathlib import Path
from openpyxl import load_workbook
import re

roots=[Path(r'C:\Users\omarf\Downloads'),Path(r'C:\Users\omarf\OneDrive\Documents\ChatGPT\Logistic ERP\reports')]
for root in roots:
    for p in root.rglob('*.xlsx'):
        if p.stat().st_size>5_000_000: continue
        try:
            wb=load_workbook(p,read_only=True,data_only=True)
            hits=[]
            for ws in wb:
                terms=[]
                for row in ws.iter_rows(min_row=1,max_row=min(8,ws.max_row),max_col=min(40,ws.max_column),values_only=True):
                    for v in row:
                        if isinstance(v,str) and re.search(r'iqama|إقامة|اقامة|الهوية|identity',v,re.I):
                            terms.append(v.strip()[:35])
                if terms: hits.append((ws.title,ws.max_row,terms[:4]))
            if hits: print(p.name, p.stat().st_mtime, hits)
        except Exception as exc: print('ERROR',p.name,str(exc)[:80])
