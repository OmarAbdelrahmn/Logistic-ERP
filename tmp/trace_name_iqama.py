from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict
b=Path(r'C:\Users\omarf\Downloads');m=load_workbook(b/'دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx',read_only=True,data_only=True);r=load_workbook(b/'المقيمين_ERP_مكتمل_السجلات_الثلاثة_12-09-2026.xlsx',read_only=True,data_only=True)
name=m['ذا شيفز - The Chefsz']['B31'].value
ids=[str(m['ذا شيفز - The Chefsz']['A31'].value)]+[str(x[0]) for x in r.active.values if x[1]==name]
for p in b.glob('*.xlsx'):
 if p.name.startswith('~$') or p.stat().st_size>5_000_000:continue
 try:
  w=load_workbook(p,read_only=True,data_only=True)
  for s in w:
   for row in s.iter_rows():
    found=[i for i,v in enumerate(row) if v.value is not None and str(v.value).strip().removesuffix('.0') in ids]
    if found:print(p.name,s.title,row[0].row,found,'idvariant',[ids.index(str(row[i].value).strip().removesuffix('.0')) for i in found])
 except Exception:pass
