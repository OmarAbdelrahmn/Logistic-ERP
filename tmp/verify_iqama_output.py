from pathlib import Path
from openpyxl import load_workbook
from collections import Counter
b=Path(r'C:\Users\omarf\Downloads\دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx')
o=Path(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx')
a=load_workbook(b,data_only=False);z=load_workbook(o,data_only=False)
print('sheets_equal',a.sheetnames==z.sheetnames)
for x,y in zip(a,z):
 changes=[]; style_changes=[]
 for row in x:
  for c in row:
   d=y[c.coordinate]
   if c.value!=d.value:changes.append(c.coordinate)
   if c._style!=d._style:style_changes.append(c.coordinate)
 print(x.title,'dimensions',x.max_row,x.max_column,y.max_row,y.max_column,'value_changes',changes,'style_changes_count',len(style_changes),'style_examples',style_changes[:8], 'merges',len(x.merged_cells.ranges),len(y.merged_cells.ranges),'cf',len(x.conditional_formatting),len(y.conditional_formatting),'tables',len(x.tables),len(y.tables),'validations',len(x.data_validations.dataValidation),len(y.data_validations.dataValidation),'freeze',x.freeze_panes,y.freeze_panes)
 print('rowheight_changes',sum(x.row_dimensions[i].height!=y.row_dimensions[i].height for i in range(1,max(x.max_row,y.max_row)+1)),'colwidth_changes',sum(x.column_dimensions[i].width!=y.column_dimensions[i].width for i in x.column_dimensions))
print('defined_names',len(a.defined_names),len(z.defined_names))
