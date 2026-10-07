from pathlib import Path
from collections import Counter,defaultdict
from openpyxl import load_workbook
import re,unicodedata
from difflib import SequenceMatcher

b=Path(r'C:\Users\omarf\Downloads')
m=load_workbook(b/'دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx',read_only=True,data_only=True)
r=load_workbook(b/'المقيمين_ERP_مكتمل_السجلات_الثلاثة_12-09-2026.xlsx',read_only=True,data_only=True)
def id(v):return re.sub(r'\D','',str(v or '').strip().removesuffix('.0'))
def name(v):
    s=unicodedata.normalize('NFKD',str(v or '').lower())
    s=''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^\w]','',s).replace('أ','ا').replace('إ','ا').replace('آ','ا').replace('ى','ي').replace('ة','ه')
def mask(x):return x[:3]+'…'+x[-2:] if len(x)>5 else x
ref={id(row[0]):(n+2,str(row[1] or '')) for n,row in enumerate(list(r.active.values)[1:]) if id(row[0])}
name_map=defaultdict(list)
for i,(row,n) in enumerate(ref.items()):name_map[name(n[1])].append((row,n[0]))
for s in list(m)[1:]:
    start=5 if s.title.startswith('كيتا') else 2
    icol=4 if s.title.startswith('كيتا') else 0 if s.title.startswith('ذا شيفز') else 1
    issues=defaultdict(list); data=[]
    for row in s.iter_rows(min_row=start):
        if all(c.value is None for c in row):continue
        n=row[0].row; v=[c.value for c in row]; iq=id(v[icol]);data.append((n,v,iq))
        if iq not in ref:issues['iqama_not_resident_master'].append(n)
        for j,x in enumerate(v):
            if x is None or str(x).strip()=='':issues[f'blank_{j+1}'].append(n)
    print('\n',s.title,'data rows',len(data))
    for k,rows in issues.items(): print(k,len(rows),rows[:25])
    if s.title.startswith('ذا شيفز'):
        print('unmatched name candidates:')
        for n,v,iq in data:
            if iq in ref: continue
            nam=name(v[1]); exact=name_map.get(nam,[])
            ranked=sorted(((SequenceMatcher(None,nam,k).ratio(),k,entries) for k,entries in name_map.items()),reverse=True)[:2]
            print(n,mask(iq),'name exact',[(mask(x),rn) for x,rn in exact], 'near',[(round(score,2),[(mask(x),rn) for x,rn in entries]) for score,k,entries in ranked if score>=.72])
    if s.title.startswith('كيتا'):
        mismatch=[]
        for n,v,iq in data:
            actual=name(str(v[2] or '')+str(v[3] or ''))
            expected=name(ref[iq][1]) if iq in ref else ''
            if actual and expected and actual not in expected and expected not in actual:
                mismatch.append((n,mask(iq),round(SequenceMatcher(None,actual,expected).ratio(),2)))
        print('name mismatch',len(mismatch),mismatch[:25])
    byacc=defaultdict(list)
    for n,v,iq in data:byacc[str(v[0])].append((n,v,iq))
    dups=[(a,items) for a,items in byacc.items() if len(items)>1]
    if dups:
        print('duplicate accounts',len(dups))
        for a,items in dups:
            print(a,[(n,mask(iq),[str(x)[:22] for x in v[2:6]]) for n,v,iq in items])
