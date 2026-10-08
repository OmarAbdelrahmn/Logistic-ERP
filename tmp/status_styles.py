from collections import defaultdict,Counter
from openpyxl import load_workbook
w=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx')
for title in ['جاهز - Jahez','هنقرستيشن - HungerStation']:
 s=w[title];d=defaultdict(Counter)
 for row in s.iter_rows(min_row=2):
  c=row[5]
  if c.value is not None:d[str(c.value)][c.style_id]+=1
 print(title,dict(d))
