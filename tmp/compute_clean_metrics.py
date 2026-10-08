from collections import Counter
from openpyxl import load_workbook

w=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx',read_only=True,data_only=True)
n=load_workbook(r'C:\Users\omarf\Downloads\تمام تطبيق هنقرستيشن.xlsx',read_only=True,data_only=True).active
data={}
for s in list(w)[1:]:
 if s.title.startswith('كيتا'):
  rows=[r for r in list(s.values)[4:] if r[0] is not None]
  data['Keeta']=[(str(r[4]),str(r[0]),r[7],r[15],r[16],r[8]) for r in rows]
 elif s.title.startswith('جاهز'):
  rows=[r for r in list(s.values)[1:] if r[0] is not None]
  seen=set();kept=[]
  for r in rows:
   if str(r[0]) in seen:continue
   seen.add(str(r[0]));kept.append(r)
  data['Jahez']=[(str(r[1]),str(r[0]),r[2],r[3],r[4],r[5]) for r in kept]
 elif s.title.startswith('ذا شيفز'):
  rows=[r for r in list(s.values)[1:] if r[0] is not None]
  data['Chefsz']=[(str(r[0]),str(r[0]),r[2],r[3],r[4],r[5]) for r in rows]
map_status={'يعمل':'يعمل (نشط)','مشاكل':'موقوف (مشاكل نظامية)','ارشيف':'غير نشط (أرشيف)'}
map_sponsor={'اكسبرس':'شركة اكسبرس جايت','البوابا':'شركة البوابا الموكبلا','التجارية':'مؤسسة البوابة التجارية','خارج الكفالة':'خارج الكفالة'}
data['HungerStation']=[(str(r[0]),str(r[2]),r[4],r[3],map_sponsor[r[5]],map_status[r[6]]) for r in list(n.values)[1:] if r[2] is not None]
for k,rows in data.items():
 print(k,'count',len(rows),'vehicles',Counter(r[2] for r in rows),'city',Counter(r[3] for r in rows),'sponsor',Counter(r[4] for r in rows),'status',Counter(r[5] for r in rows))
print('total',sum(len(v) for v in data.values()),'iqamas',len(set(r[0] for rows in data.values() for r in rows)))
