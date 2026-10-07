from pathlib import Path
from collections import defaultdict,Counter
from openpyxl import load_workbook
import re,unicodedata
b=Path(r'C:\Users\omarf\Downloads'); m=load_workbook(b/'دليل_معرفات_وحسابات_التطبيقات_الموحد_2026-10-07.xlsx',read_only=True,data_only=True);r=load_workbook(b/'المقيمين_ERP_مكتمل_السجلات_الثلاثة_12-09-2026.xlsx',read_only=True,data_only=True)
def clean(v):return re.sub(r'\D','',str(v or '').removesuffix('.0'))
def mask(v):return v[:3]+'…'+v[-2:]
ref={clean(row[0]):(n+2,row[1]) for n,row in enumerate(list(r.active.values)[1:]) if clean(row[0])}
un=defaultdict(list)
for s in list(m)[1:]:
 start=5 if s.title.startswith('كيتا') else 2
 col=4 if s.title.startswith('كيتا') else 0 if s.title.startswith('ذا شيفز') else 1
 for row in s.iter_rows(min_row=start):
  x=[c.value for c in row];i=clean(x[col])
  if not i:continue
  if i not in ref:un[i].append((s.title,row[0].row,x))
print('unmatched unique',len(un),'on both sheets',sum(len(v)>1 for v in un.values()))
for iq,rows in un.items():
 print(mask(iq),[(s,n,str(x[-2])[:35],str(x[-1])[:35]) for s,n,x in rows])
print('chefsz row 31')
c=m['ذا شيفز - The Chefsz']; vals=[x.value for x in c[31]]
print('main row31',mask(clean(vals[0])),str(vals[1]),vals[2:])
for no,row in enumerate(r.active.iter_rows(min_row=2,values_only=True),2):
 if str(row[1]).strip()==str(vals[1]).strip(): print('resident same name',no,mask(clean(row[0])),str(row[1]),row[14:16])
