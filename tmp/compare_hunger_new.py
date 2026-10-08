from pathlib import Path
from collections import Counter,defaultdict
from openpyxl import load_workbook
import re

old=load_workbook(r'outputs\iqama-platform-review-20261007\platform_accounts_iqama_corrected.xlsx',read_only=True,data_only=True)['هنقرستيشن - HungerStation']
new=load_workbook(r'C:\Users\omarf\Downloads\تمام تطبيق هنقرستيشن.xlsx',read_only=True,data_only=True).active
def st(v):return str(v or '').strip().removesuffix('.0')
def mask(v):
 s=st(v);return s[:3]+'…'+s[-2:] if len(s)>5 else s
def read(ws,newfile):
 d=defaultdict(list);issues=Counter();statuses=Counter();sponsors=Counter();data=[]
 for n,row in enumerate(list(ws.values)[1:],2):
  if not any(v is not None for v in row):continue
  account=st(row[2] if newfile else row[0]);iqama=st(row[0] if newfile else row[1]);status=st(row[6] if newfile else row[5]);city=st(row[3]);vehicle=st(row[4] if newfile else row[2]);sponsor=st(row[5] if newfile else row[4]);
  if not account:issues['blank account']+=1
  if not iqama:issues['blank iqama']+=1
  if iqama and not re.fullmatch(r'[12]\d{9}',iqama):issues['invalid iqama']+=1
  d[account].append((n,iqama,status,city,vehicle,sponsor));data.append((n,account,iqama,status,city,vehicle,sponsor));statuses[status]+=1;sponsors[sponsor]+=1
 return d,data,issues,statuses,sponsors
od,rows_old,oi,os,osp=read(old,False)
nd,rows_new,ni,ns,nsp=read(new,True)
print('old',len(rows_old),'accounts',len(od),'issues',oi,'statuses',os,'sponsors',osp)
print('new',len(rows_new),'accounts',len(nd),'issues',ni,'statuses',ns,'sponsors',nsp)
print('old duplicate accounts',sum(len(x)>1 for x in od.values()),'new duplicate accounts',sum(len(x)>1 for x in nd.values()))
print('old-only accounts',len(od.keys()-nd.keys()),[(a,[(n,mask(i),s) for n,i,s,*_ in od[a]]) for a in sorted(od.keys()-nd.keys())][:40])
print('new-only accounts',len(nd.keys()-od.keys()),[(a,[(n,mask(i),s) for n,i,s,*_ in nd[a]]) for a in sorted(nd.keys()-od.keys())][:40])
changed=[]
for a in od.keys()&nd.keys():
 oldids={x[1] for x in od[a]};newids={x[1] for x in nd[a]}
 if oldids!=newids:changed.append((a,[(n,mask(i)) for n,i,*_ in od[a]],[(n,mask(i)) for n,i,*_ in nd[a]]))
print('iqama changed',len(changed),changed[:80])
statusdiff=[]
for a in od.keys()&nd.keys():
 oldstatuses={x[2] for x in od[a]};newstatuses={x[2] for x in nd[a]}
 if oldstatuses!=newstatuses:statusdiff.append((a,oldstatuses,newstatuses))
print('status changed count',len(statusdiff),'examples',statusdiff[:15])
