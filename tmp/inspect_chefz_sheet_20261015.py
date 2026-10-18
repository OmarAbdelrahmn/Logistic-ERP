import openpyxl, json, datetime, hashlib
from pathlib import Path
p=Path(r'C:\Users\omarf\Downloads\داش شيفز  حسبات سيستم مطلوب بتاريخ 14-10-2026.xlsx')
w=openpyxl.load_workbook(p,data_only=True)
out={'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheets':[]}
def val(v):
    return v.isoformat() if isinstance(v,(datetime.date,datetime.datetime)) else v
for s in w:
    rows=[{'row':r[0].row,'cells':{c.column_letter:val(c.value) for c in r if c.value is not None}} for r in s if any(c.value is not None for c in r)]
    out['sheets'].append({'name':s.title,'maxRow':s.max_row,'maxCol':s.max_column,'rows':rows})
    print(json.dumps({'sheet':s.title,'rows':s.max_row,'cols':s.max_column},ensure_ascii=False))
private=Path('artifacts/chefz-import-20261015');private.mkdir(parents=True,exist_ok=True)
(private/'source.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
