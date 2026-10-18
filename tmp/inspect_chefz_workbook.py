import openpyxl, json
from pathlib import Path
p=Path(r'C:\Users\omarf\Downloads\داش شيفز مطلوووووووووب222.xlsx')
w=openpyxl.load_workbook(p, data_only=False)
data=[]
for s in w:
    rows=[]
    for row in s:
        cells={c.coordinate:str(c.value) for c in row if c.value is not None}
        if cells: rows.append(cells)
    data.append({'sheet':s.title,'rows':s.max_row,'cols':s.max_column,'merges':[str(m) for m in s.merged_cells.ranges],'cells':rows})
Path('tmp/chefz-workbook.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
for s in data:
    print(json.dumps({**s,'cells':s['cells'][:18]},ensure_ascii=False))
